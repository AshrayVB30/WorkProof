from beanie import PydanticObjectId
from datetime import date, datetime, timedelta
from typing import List, Optional
from app.models.activity import Activity, WorkSession, DailyReport, ActivityType
from app.models.user import User
from io import BytesIO
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


async def generate_daily_report(user_id: PydanticObjectId, report_date: date) -> Optional[DailyReport]:
    """Generate or update a daily report for a user"""
    # Get all activities for the day
    start_datetime = datetime.combine(report_date, datetime.min.time())
    end_datetime = datetime.combine(report_date, datetime.max.time())
    
    activities = await Activity.find(
        Activity.user_id == user_id,
        Activity.timestamp >= start_datetime,
        Activity.timestamp <= end_datetime
    ).to_list()
    
    # If no activities, return None
    if not activities:
        return None
    
    # Calculate metrics
    total_activities = len(activities)
    active_seconds = sum(a.duration_seconds for a in activities if a.activity_type == ActivityType.ACTIVE)
    idle_seconds = sum(a.duration_seconds for a in activities if a.activity_type == ActivityType.IDLE)
    total_seconds = active_seconds + idle_seconds
    
    total_hours = total_seconds / 3600
    active_hours = active_seconds / 3600
    idle_hours = idle_seconds / 3600
    productive_percentage = (active_seconds / total_seconds * 100) if total_seconds > 0 else 0
    
    # Application usage breakdown
    app_usage = {}
    for activity in activities:
        if activity.application_name:
            app_usage[activity.application_name] = app_usage.get(activity.application_name, 0) + activity.duration_seconds
    
    # Count screenshots
    screenshot_count = sum(1 for a in activities if a.screenshot_path)
    
    # Get work sessions for the day
    sessions = await WorkSession.find(
        WorkSession.user_id == user_id,
        WorkSession.start_time >= start_datetime,
        WorkSession.start_time <= end_datetime
    ).to_list()
    
    session_count = len(sessions)
    
    # Check if report already exists
    existing_report = await DailyReport.find_one(
        DailyReport.user_id == user_id,
        DailyReport.report_date == report_date
    )
    
    if existing_report:
        # Update existing report
        existing_report.total_hours = round(total_hours, 2)
        existing_report.active_hours = round(active_hours, 2)
        existing_report.idle_hours = round(idle_hours, 2)
        existing_report.productive_percentage = round(productive_percentage, 2)
        existing_report.total_activities = total_activities
        existing_report.session_count = session_count
        existing_report.screenshot_count = screenshot_count
        existing_report.app_usage = app_usage
        existing_report.generated_at = datetime.utcnow()
        
        await existing_report.save()
        return existing_report
    else:
        # Create new report
        new_report = DailyReport(
            user_id=user_id,
            report_date=report_date,
            total_hours=round(total_hours, 2),
            active_hours=round(active_hours, 2),
            idle_hours=round(idle_hours, 2),
            productive_percentage=round(productive_percentage, 2),
            total_activities=total_activities,
            session_count=session_count,
            screenshot_count=screenshot_count,
            app_usage=app_usage
        )
        
        await new_report.insert()
        return new_report


def generate_pdf_report(user: User, reports: List[DailyReport], start_date: date, end_date: date) -> bytes:
    """Generate a PDF report"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    
    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=30,
        alignment=1  # Center
    )
    
    elements.append(Paragraph("WorkProof Activity Report", title_style))
    elements.append(Spacer(1, 0.2 * inch))
    
    # User info
    user_info = f"""
    <b>Employee:</b> {user.full_name}<br/>
    <b>Employee ID:</b> {user.employee_id or 'N/A'}<br/>
    <b>Department:</b> {user.department or 'N/A'}<br/>
    <b>Report Period:</b> {start_date} to {end_date}<br/>
    <b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    """
    elements.append(Paragraph(user_info, styles['Normal']))
    elements.append(Spacer(1, 0.3 * inch))
    
    # Summary table
    if reports:
        total_hours = sum(r.total_hours for r in reports)
        total_active = sum(r.active_hours for r in reports)
        avg_productive = sum(r.productive_percentage for r in reports) / len(reports)
        
        summary_data = [
            ['Metric', 'Value'],
            ['Total Days', str(len(reports))],
            ['Total Hours Worked', f'{total_hours:.2f} hrs'],
            ['Active Hours', f'{total_active:.2f} hrs'],
            ['Average Productivity', f'{avg_productive:.1f}%'],
            ['Expected Hours', f'{user.expected_daily_hours * len(reports):.2f} hrs']
        ]
        
        summary_table = Table(summary_data, colWidths=[3 * inch, 3 * inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4a90e2')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(Paragraph("<b>Summary</b>", styles['Heading2']))
        elements.append(Spacer(1, 0.1 * inch))
        elements.append(summary_table)
        elements.append(Spacer(1, 0.3 * inch))
        
        # Daily breakdown
        elements.append(Paragraph("<b>Daily Breakdown</b>", styles['Heading2']))
        elements.append(Spacer(1, 0.1 * inch))
        
        daily_data = [['Date', 'Total Hours', 'Active Hours', 'Productivity %', 'Activities']]
        for report in reports:
            daily_data.append([
                str(report.report_date),
                f'{report.total_hours:.2f}',
                f'{report.active_hours:.2f}',
                f'{report.productive_percentage:.1f}%',
                str(report.total_activities)
            ])
        
        daily_table = Table(daily_data, colWidths=[1.5*inch, 1.2*inch, 1.2*inch, 1.2*inch, 1*inch])
        daily_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4a90e2')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(daily_table)
    
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(user: User, reports: List[DailyReport], start_date: date, end_date: date) -> bytes:
    """Generate an Excel report"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Activity Report"
    
    # Header styling
    header_fill = PatternFill(start_color="4A90E2", end_color="4A90E2", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Title
    ws['A1'] = "WorkProof Activity Report"
    ws['A1'].font = Font(bold=True, size=16)
    ws.merge_cells('A1:E1')
    
    # User info
    ws['A3'] = "Employee:"
    ws['B3'] = user.full_name
    ws['A4'] = "Employee ID:"
    ws['B4'] = user.employee_id or 'N/A'
    ws['A5'] = "Department:"
    ws['B5'] = user.department or 'N/A'
    ws['A6'] = "Report Period:"
    ws['B6'] = f"{start_date} to {end_date}"
    
    # Summary
    if reports:
        total_hours = sum(r.total_hours for r in reports)
        total_active = sum(r.active_hours for r in reports)
        avg_productive = sum(r.productive_percentage for r in reports) / len(reports)
        
        ws['A8'] = "Summary"
        ws['A8'].font = Font(bold=True, size=14)
        
        ws['A9'] = "Total Days"
        ws['B9'] = len(reports)
        ws['A10'] = "Total Hours Worked"
        ws['B10'] = round(total_hours, 2)
        ws['A11'] = "Active Hours"
        ws['B11'] = round(total_active, 2)
        ws['A12'] = "Average Productivity"
        ws['B12'] = f"{avg_productive:.1f}%"
        
        # Daily breakdown
        ws['A14'] = "Daily Breakdown"
        ws['A14'].font = Font(bold=True, size=14)
        
        headers = ['Date', 'Total Hours', 'Active Hours', 'Idle Hours', 'Productivity %', 'Activities', 'Sessions']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=15, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Data rows
        for row, report in enumerate(reports, start=16):
            ws.cell(row=row, column=1, value=str(report.report_date))
            ws.cell(row=row, column=2, value=round(report.total_hours, 2))
            ws.cell(row=row, column=3, value=round(report.active_hours, 2))
            ws.cell(row=row, column=4, value=round(report.idle_hours, 2))
            ws.cell(row=row, column=5, value=f"{report.productive_percentage:.1f}%")
            ws.cell(row=row, column=6, value=report.total_activities)
            ws.cell(row=row, column=7, value=report.session_count)
    
    # Adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
